#pip install python-docx

from docx import Document

doc = Document()

doc.add_heading("Test case report")
doc.add_paragraph("Test case executed successfully")
doc.add_paragraph("Status : Passed")
doc.save("TestReport.doc")

# excel
#pip install openpyxl
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws['A1'] = "SDET Training"
wb.save("excel_demo_1.xlsx")

#text

with open("C:\python-selenium-sdet7\write_file1.txt","w") as fw:
    fw.write("first line")

with open("C:\python-selenium-sdet7\write_file1.txt","r") as fr:
    print(fr.read())

f = open("C:\python-selenium-sdet7\write_file.txt","w")
list1 = [10,20,30,40]
for items in list1:
    f.write(str(items) + "\n")
#f.write("this is my first automated file")
f.close()

f = open("C:\python-selenium-sdet7\write_file.txt","r+")
#print (f.read())
f.write("read write text")
f.close()

#to read entire excel
from openpyxl import Workbook , load_workbook
wb = load_workbook(filename="excel_demo_1.xlsx")
#sh = wb.active
sh =wb['Sheet']
print (sh['A2'].value)
print (sh.cell(row=2,column=2).value) # particular cell

row_ct = sh.max_row
col_ct = sh.max_column

for i in range(1,row_ct+1):
    for j in range(1,col_ct+1):
        print (sh.cell(row=i, column =j).value, end =' ')  #reading entire excel
    print ('\n')

# pip install Faker

from faker import Faker
fake_data = Faker()
print (fake_data.name())
print (fake_data.email())
print (fake_data.address())
from faker import Faker
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
fake_data = Faker()

for i in range(1,100):
    for j in range(1,4):
        ws.cell(row=i,column =1).value = fake_data.name()
        ws.cell(row=i, column=2).value = fake_data.email()
        ws.cell(row=i, column=3).value = fake_data.address()
wb.save("excel_test_data7.xlsx")

#To create multiple columns
from faker import Faker
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
fake_data = Faker()

for i in range(1,100):
    for j in range(1,4):
        ws.cell(row=i,column =1).value = fake_data.name()
        ws.cell(row=i, column=2).value = fake_data.email()
        ws.cell(row=i, column=3).value = fake_data.address()
wb.save("excel_test_data7.xlsx")


